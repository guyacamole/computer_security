import os
from docx import Document
from openpyxl import load_workbook
from PyPDF2 import PdfReader


def extract_metadata(file_path):
    file_name, file_ext = os.path.splitext(file_path)
    file_ext = file_ext.lower()

    if file_ext == '.docx':
        doc = Document(file_path)
        metadata = doc.core_properties
        print(f"Metadata for {file_path}:")
        print(f"Title: {metadata.title}")
        print(f"Author: {metadata.author}")
        print(f"Subject: {metadata.subject}")
        print(f"Keywords: {metadata.keywords}")
        print(f"Created: {metadata.created}")
        print(f"Modified: {metadata.modified}")
        print()

    elif file_ext == '.xlsx':
        wb = load_workbook(file_path)
        metadata = wb.properties
        print(f"Metadata for {file_path}:")
        print(f"Title: {metadata.title}")
        print(f"Author: {metadata.creator}")
        print(f"Subject: {metadata.subject}")
        print(f"Keywords: {metadata.keywords}")
        print(f"Created: {metadata.created}")
        print(f"Modified: {metadata.modified}")
        print()

    elif file_ext == '.pdf':
        with open(file_path, 'rb') as file:
            pdf = PdfReader(file)
            metadata = pdf.metadata
            print(f"Metadata for {file_path}:")
            print(f"Title: {metadata.title}")
            print(f"Author: {metadata.author}")
            print(f"Subject: {metadata.subject}")
            print(f"Creator: {metadata.creator}")
            print(f"Producer: {metadata.producer}")
            print()

    else:
        print(f"Unsupported file format: {file_ext}")


def main():
    directory = input("Enter the directory path: ")
    for root, dirs, files in os.walk(directory):
        for file in files:
            file_path = os.path.join(root, file)
            extract_metadata(file_path)


if __name__ == "__main__":
    main()
